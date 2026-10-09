"""Builds the Cleaning Business Job Profit & Quote Calculator workbooks.

Usage: python3 build.py            -> writes dist/*.xlsx (3 tiers x blank/example)

No macros. Only functions that work in both Excel and Google Sheets
(IF, IFERROR, SUM, SUMIFS, COUNT, COUNTIF, MAX, ROUNDUP, EDATE, DATE, TEXT).
"""
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dist")

NAVY = "1F3A4D"
TEAL = "2A7F8E"
INPUT_FILL = PatternFill("solid", fgColor="FFF6D6")
CALC_FILL = PatternFill("solid", fgColor="EEF2F4")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
KEY_FILL = PatternFill("solid", fgColor="D8EEF1")
BAD_FILL = PatternFill("solid", fgColor="F8D0D0")
GOOD_FILL = PatternFill("solid", fgColor="D5EFD5")
thin = Side(style="thin", color="B8C2C8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '"$"#,##0.00'
PCT = "0.0%"
HRS = "0.00"

DISCLAIMER = ("This workbook is a planning tool. It is not tax, accounting, legal or "
              "financial advice. Check your numbers with your own accountant.")


def title(ws, text, sub=None, width=6):
    ws["A1"] = text
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="555555")
    ws.sheet_view.showGridLines = False


def header_row(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def style_input(c, fmt=None):
    c.fill = INPUT_FILL
    c.border = BOX
    if fmt:
        c.number_format = fmt


def style_calc(c, fmt=None, key=False):
    c.fill = KEY_FILL if key else CALC_FILL
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if key:
        c.font = Font(bold=True)


# ---------------------------------------------------------------- Settings
SET_ROWS = {
    "labor": 4, "overhead": 5, "margin": 6, "mile": 7, "disc": 8, "minprice": 9,
}


def build_settings(ws, example):
    title(ws, "Settings", "Starter values are filled in so the Examples work. REPLACE them with your own numbers.")
    header_row(ws, 3, ["Setting", "Value", "What it means"])
    rows = [
        ("Labor cost per cleaner-hour", 18, USD,
         "What one cleaner costs you per hour: wages PLUS payroll taxes/workers' comp. Pay yourself too."),
        ("Overhead (% added on top of direct cost)", 0.10, PCT,
         "Insurance, software, phone, cleaning equipment wear, admin time. 8-15% is common; use your real number."),
        ("Target profit margin", 0.35, PCT,
         "Profit as a % of the PRICE (not of cost). 35% margin means $35 of every $100 you charge is profit."),
        ("Vehicle cost per mile", 0.67, USD,
         "Your real cost to drive one mile (fuel, wear, insurance). Check the current IRS mileage rate as a reference."),
        ("Recurring-client discount", 0.10, PCT,
         "Discount off the one-time price for recurring clients. Enter 0% for none."),
        ("Minimum job price", 90, USD,
         "Never quote below this, even for a tiny job."),
    ]
    for i, (lab, val, fmt, note) in enumerate(rows):
        r = 4 + i
        ws.cell(r, 1, lab).border = BOX
        ws.cell(r, 1).alignment = Alignment(vertical="center", wrap_text=True)
        c = ws.cell(r, 2, val)
        c.alignment = Alignment(vertical="center")
        style_input(c, fmt)
        n = ws.cell(r, 3, note)
        n.alignment = Alignment(wrap_text=True, vertical="top")
        n.border = BOX
        ws.row_dimensions[r].height = 45
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 80
    ws["A11"] = DISCLAIMER
    ws["A11"].font = Font(italic=True, color="777777")
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="0.95",
                        showErrorMessage=True, errorTitle="Margin",
                        error="Enter a percentage between 0% and 95%.")
    ws.add_data_validation(dv)
    dv.add("B6")
    dv2 = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                         showErrorMessage=True, error="Enter a percentage between 0% and 100%.")
    ws.add_data_validation(dv2)
    dv2.add("B5")
    dv2.add("B8")
    dv3 = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                         showErrorMessage=True, error="Enter a number of 0 or more.")
    ws.add_data_validation(dv3)
    dv3.add("B4")
    dv3.add("B7")
    dv3.add("B9")


S_LABOR, S_OVER, S_MARGIN, S_MILE, S_DISC, S_MIN = (
    "Settings!$B$4", "Settings!$B$5", "Settings!$B$6", "Settings!$B$7", "Settings!$B$8", "Settings!$B$9")

# ----------------------------------------------------------- Quote block
# Row layout shared by Quote sheet (col B) and Examples sheet (cols B..D)
Q = {"name": 3, "type": 4, "cleaners": 5, "hours": 6, "materials": 7, "miles": 8,
     "other": 9, "visits": 10,
     "tot_hours": 13, "labor": 14, "travel": 15, "direct": 16, "overhead": 17, "cost": 18,
     "price_full": 19, "disc": 20, "price": 21, "profit": 22, "margin": 23, "markup": 24,
     "pph": 25, "m_rev": 26, "m_profit": 27, "flag": 28}

INPUT_LABELS = [
    (Q["name"], "Client / job name", None),
    (Q["type"], "Job type (One-time or Recurring)", None),
    (Q["cleaners"], "Number of cleaners", "0"),
    (Q["hours"], "Hours on site per cleaner", HRS),
    (Q["materials"], "Supplies & materials for this visit ($)", USD),
    (Q["miles"], "Round-trip miles (all vehicles)", "0.0"),
    (Q["other"], "Other direct costs ($, e.g. equipment rental)", USD),
    (Q["visits"], "Visits per month (recurring jobs only)", "0"),
]
CALC_LABELS = [
    (Q["tot_hours"], "Total labor hours", HRS, False),
    (Q["labor"], "Labor cost", USD, False),
    (Q["travel"], "Travel cost", USD, False),
    (Q["direct"], "Direct cost (labor + travel + supplies + other)", USD, False),
    (Q["overhead"], "Overhead", USD, False),
    (Q["cost"], "TOTAL COST to do this visit", USD, True),
    (Q["price_full"], "Price needed to hit your target margin", USD, False),
    (Q["disc"], "Recurring discount applied", USD, False),
    (Q["price"], "QUOTE THIS PRICE (per visit)", USD, True),
    (Q["profit"], "Profit per visit", USD, True),
    (Q["margin"], "Profit margin (profit / price)", PCT, True),
    (Q["markup"], "Markup on cost (profit / cost)", PCT, False),
    (Q["pph"], "Profit per labor hour", USD, False),
    (Q["m_rev"], "Monthly revenue (recurring)", USD, False),
    (Q["m_profit"], "Monthly profit (recurring)", USD, False),
    (Q["flag"], "Check", None, False),
]


def calc_formulas(c):
    """Return {row: formula} for column letter c."""
    g = lambda k: f"{c}{Q[k]}"
    return {
        Q["tot_hours"]: f"={g('cleaners')}*{g('hours')}",
        Q["labor"]: f"={g('tot_hours')}*{S_LABOR}",
        Q["travel"]: f"={g('miles')}*{S_MILE}",
        Q["direct"]: f"={g('labor')}+{g('travel')}+{g('materials')}+{g('other')}",
        Q["overhead"]: f"={g('direct')}*{S_OVER}",
        Q["cost"]: f"={g('direct')}+{g('overhead')}",
        Q["price_full"]: f'=IF({S_MARGIN}>=1,0,{g("cost")}/(1-{S_MARGIN}))',
        Q["disc"]: f'=IF({g("type")}="Recurring",{g("price_full")}*{S_DISC},0)',
        Q["price"]: f"=MAX({S_MIN},ROUNDUP({g('price_full')}-{g('disc')},0))",
        Q["profit"]: f"={g('price')}-{g('cost')}",
        Q["margin"]: f"=IF({g('price')}=0,0,{g('profit')}/{g('price')})",
        Q["markup"]: f"=IF({g('cost')}=0,0,{g('profit')}/{g('cost')})",
        Q["pph"]: f"=IF({g('tot_hours')}=0,0,{g('profit')}/{g('tot_hours')})",
        Q["m_rev"]: f'=IF({g("type")}="Recurring",{g("price")}*{g("visits")},0)',
        Q["m_profit"]: f'=IF({g("type")}="Recurring",{g("profit")}*{g("visits")},0)',
        Q["flag"]: (f'=IF({S_MARGIN}>=1,"Fix target margin on Settings (must be under 100%)",'
                    f'IF({g("tot_hours")}=0,"Enter cleaners and hours to get a result",'
                    f'IF({g("profit")}<0,"LOSING MONEY - raise the price",'
                    f'IF({g("margin")}<{S_MARGIN}-0.0001,"Below target margin because of the recurring discount - fine if intended",'
                    f'"OK - at or above target margin"))))'),
    }


def add_flag_cf(ws, ref):
    first = ref.split(":")[0]
    amber = PatternFill("solid", fgColor="FBE7B5")
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({first},2)="OK"'], fill=GOOD_FILL))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({first},5)="Below"'], fill=amber))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'AND(LEFT({first},2)<>"OK",LEFT({first},5)<>"Below")'], fill=BAD_FILL))


def write_block(ws, cols, jobs):
    """Write labels in col A, inputs+calcs in each column of cols."""
    for r, lab, fmt in INPUT_LABELS:
        ws.cell(r, 1, lab).border = BOX
    for r, lab, fmt, key in CALC_LABELS:
        c = ws.cell(r, 1, lab)
        c.border = BOX
        if key:
            c.font = Font(bold=True)
    ws.cell(12, 1, "RESULTS").font = Font(bold=True, color=NAVY)
    for col, job in zip(cols, jobs):
        for r, lab, fmt in INPUT_LABELS:
            key = [k for k, v in Q.items() if v == r][0]
            cell = ws[f"{col}{r}"]
            cell.value = job.get(key) if job else None
            style_input(cell, fmt)
        f = calc_formulas(col)
        for r, lab, fmt, key in CALC_LABELS:
            cell = ws[f"{col}{r}"]
            cell.value = f[r]
            style_calc(cell, fmt, key)
            if r == Q["flag"]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[col].width = 26
        ws.row_dimensions[Q["flag"]].height = 48
    ws.column_dimensions["A"].width = 48
    add_flag_cf(ws, f"{cols[0]}{Q['flag']}:{cols[-1]}{Q['flag']}")


def add_input_validation(ws, cols):
    rng = lambda r: " ".join(f"{c}{r}" for c in cols)
    dv_type = DataValidation(type="list", formula1='"One-time,Recurring"', allow_blank=True,
                             showErrorMessage=True, error="Choose One-time or Recurring.")
    ws.add_data_validation(dv_type)
    dv_type.add(rng(Q["type"]))
    dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            showErrorMessage=True, error="Enter a number of 0 or more.")
    ws.add_data_validation(dv_num)
    for k in ("cleaners", "hours", "materials", "miles", "other", "visits"):
        dv_num.add(rng(Q[k]))


JOBS = [
    {"name": "Standard 3BR/2BA - biweekly", "type": "Recurring", "cleaners": 2, "hours": 2.0,
     "materials": 6, "miles": 14, "other": 0, "visits": 2},
    {"name": "Move-out deep clean - 4BR", "type": "One-time", "cleaners": 3, "hours": 3.0,
     "materials": 25, "miles": 22, "other": 15, "visits": None},
    {"name": "Small office - nightly, 5x/week", "type": "Recurring", "cleaners": 1, "hours": 1.5,
     "materials": 5, "miles": 8, "other": 0, "visits": 22},
]

# actual results for the example tracker (job index, date, actual hours, materials, miles, other)
ACTUALS = [
    (0, date(2026, 9, 12), 4.5, 7, 14, 0),     # ran a bit long
    (1, date(2026, 9, 19), 12.0, 30, 22, 15),  # big overrun: under-quoted hours
    (2, date(2026, 9, 22), 1.5, 5, 8, 0),      # on plan
]


def build_quote(ws, example):
    title(ws, "Job Quote Calculator",
          "Fill the yellow cells. Gray cells calculate. Quote the 'QUOTE THIS PRICE' number.")
    write_block(ws, ["B"], [JOBS[0] if example else None])
    ws.column_dimensions["C"].width = 60
    notes = {
        Q["margin"]: "Margin = profit / PRICE. This is what you keep from every dollar charged.",
        Q["markup"]: "Markup = profit / COST. A 50% markup is only a 33% margin. Don't mix them up.",
        Q["price"]: "Rounded up to the next whole dollar and never below your minimum job price.",
        Q["pph"]: "Profit left after ALL costs, divided by labor hours.",
    }
    for r, t in notes.items():
        n = ws.cell(r, 3, t)
        n.font = Font(italic=True, color="666666")
        n.alignment = Alignment(wrap_text=True, vertical="top")
    add_input_validation(ws, ["B"])
    ws["A30"] = DISCLAIMER
    ws["A30"].font = Font(italic=True, color="777777")


def build_examples(ws):
    title(ws, "Three example jobs (fictional)",
          "Same math as the Quote sheet, using your Settings. Change Settings to see prices move.")
    ws["B3"].value = None
    write_block(ws, ["B", "C", "D"], JOBS)
    ws.freeze_panes = "B3"


# ----------------------------------------------------------------- Tracker
T_FIRST, T_LAST = 5, 104


def build_tracker(ws, example):
    title(ws, "Job Tracker", "One row per completed job. Enter the quote and your ACTUAL numbers. Gray columns calculate.")
    heads = ["Date", "Client / job", "Price charged", "Estimated cost (from quote)", "Actual labor hours (total)",
             "Actual supplies ($)", "Actual miles", "Other costs ($)", "Actual total cost", "Actual profit",
             "Actual margin", "Planned profit", "Profit vs plan", "Check"]
    header_row(ws, 4, heads)
    ws.row_dimensions[4].height = 45
    widths = [12, 34, 14, 16, 14, 13, 11, 13, 14, 13, 11, 13, 13, 40]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    fmts = ["yyyy-mm-dd", None, USD, USD, HRS, USD, "0.0", USD]
    for r in range(T_FIRST, T_LAST + 1):
        for i, fmt in enumerate(fmts):
            style_input(ws.cell(r, 1 + i), fmt)
        f = {
            9: f'=IF(C{r}="","",(E{r}*{S_LABOR}+F{r}+G{r}*{S_MILE}+H{r})*(1+{S_OVER}))',
            10: f'=IF(C{r}="","",C{r}-I{r})',
            11: f'=IF(C{r}="","",IF(C{r}=0,0,J{r}/C{r}))',
            12: f'=IF(C{r}="","",C{r}-D{r})',
            13: f'=IF(C{r}="","",J{r}-L{r})',
            14: f'=IF(C{r}="","",IF(J{r}<0,"LOST MONEY",IF(K{r}<{S_MARGIN}-0.0001,"Below target margin","OK")))',
        }
        fm = {9: USD, 10: USD, 11: PCT, 12: USD, 13: USD, 14: None}
        for col, formula in f.items():
            c = ws.cell(r, col, formula)
            style_calc(c, fm[col])
    rng = f"N{T_FIRST}:N{T_LAST}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'N{T_FIRST}="LOST MONEY"'], fill=BAD_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'N{T_FIRST}="Below target margin"'],
                                                   fill=PatternFill("solid", fgColor="FBE7B5")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'N{T_FIRST}="OK"'], fill=GOOD_FILL))
    ws.freeze_panes = "C5"
    if example:
        for i, (ji, d, hrs, mat, miles, other) in enumerate(ACTUALS):
            r = T_FIRST + i
            col = "BCD"[ji]
            ws.cell(r, 1, d)
            ws.cell(r, 2, JOBS[ji]["name"])
            ws.cell(r, 3, f"=Examples!{col}{Q['price']}")
            ws.cell(r, 4, f"=Examples!{col}{Q['cost']}")
            ws.cell(r, 5, hrs)
            ws.cell(r, 6, mat)
            ws.cell(r, 7, miles)
            ws.cell(r, 8, other)


def build_dashboard(ws):
    title(ws, "Dashboard", "Updates automatically from the Job Tracker.")
    rng = lambda col: f"'Job Tracker'!${col}${T_FIRST}:${col}${T_LAST}"
    items = [
        ("Jobs tracked", f"=COUNT({rng('C')})", "0"),
        ("Total revenue", f"=SUM({rng('C')})", USD),
        ("Total actual cost", f"=SUM({rng('I')})", USD),
        ("Total actual profit", f"=SUM({rng('J')})", USD),
        ("Overall profit margin", "=IF(B5=0,0,B7/B5)", PCT),
        ("Average profit per job", "=IF(B4=0,0,B7/B4)", USD),
        ("Profit per labor hour", f"=IF(SUM({rng('E')})=0,0,B7/SUM({rng('E')}))", USD),
        ("Jobs that lost money", f'=COUNTIF({rng("J")},"<0")', "0"),
        ("Jobs below target margin", f'=COUNTIF({rng("K")},"<"&{S_MARGIN})', "0"),
        ("Total profit vs plan", f"=SUM({rng('M')})", USD),
    ]
    header_row(ws, 3, ["Measure", "Value"])
    for i, (lab, formula, fmt) in enumerate(items):
        r = 4 + i
        ws.cell(r, 1, lab).border = BOX
        c = ws.cell(r, 2, formula)
        style_calc(c, fmt, key=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws["A16"] = ("How to read it: if 'Profit vs plan' is negative, jobs are taking longer or costing more than "
                 "you quoted. Raise the hours/price you use for that kind of job.")
    ws["A16"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A16:D18")


def build_twelve_month(ws, example):
    title(ws, "12-Month Tracker", "Enter the year and a monthly profit goal. Revenue, cost and profit fill in from the Job Tracker.")
    ws["A3"] = "Year"
    ws["B3"] = 2026
    style_input(ws["B3"], "0")
    ws["A3"].font = Font(bold=True)
    header_row(ws, 5, ["Month", "Revenue", "Actual cost", "Profit", "Margin", "Profit goal", "Over / under goal"])
    rng = lambda col: f"'Job Tracker'!${col}${T_FIRST}:${col}${T_LAST}"
    for m in range(1, 13):
        r = 5 + m
        ws.cell(r, 1, f"=DATE($B$3,{m},1)").number_format = "mmm yyyy"
        for col, src in ((2, "C"), (3, "I"), (4, "J")):
            c = ws.cell(r, col, f'=SUMIFS({rng(src)},{rng("A")},">="&$A{r},{rng("A")},"<"&EDATE($A{r},1))')
            style_calc(c, USD)
        style_calc(ws.cell(r, 5, f"=IF(B{r}=0,0,D{r}/B{r})"), PCT)
        g = ws.cell(r, 6, 1500 if example else None)
        style_input(g, USD)
        style_calc(ws.cell(r, 7, f'=IF(F{r}="","",D{r}-F{r})'), USD)
        ws.cell(r, 1).border = BOX
    r = 18
    ws.cell(r, 1, "YEAR TOTAL").font = Font(bold=True)
    for col in "BCD":
        style_calc(ws[f"{col}{r}"], USD, True)
        ws[f"{col}{r}"] = f"=SUM({col}6:{col}17)"
    style_calc(ws[f"E{r}"], PCT, True)
    ws[f"E{r}"] = f"=IF(B{r}=0,0,D{r}/B{r})"
    style_calc(ws[f"F{r}"], USD, True)
    ws[f"F{r}"] = "=SUM(F6:F17)"
    style_calc(ws[f"G{r}"], USD, True)
    ws[f"G{r}"] = f"=D{r}-F{r}"
    ws.conditional_formatting.add("G6:G18", CellIsRule(operator="lessThan", formula=["0"], fill=BAD_FILL))
    ws.conditional_formatting.add("G6:G18", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=GOOD_FILL))
    ws.column_dimensions["A"].width = 16
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 16


def build_clients(ws, example):
    title(ws, "Client Tracker", "One row per client. Sort or filter by any column.")
    heads = ["Client name", "Phone", "Email", "Address / area", "Service type", "Frequency", "Price per visit",
             "Start date", "Last service", "Next follow-up", "Status", "Notes"]
    header_row(ws, 4, heads)
    for i, w in enumerate([24, 15, 26, 28, 20, 14, 14, 12, 12, 14, 12, 40]):
        ws.column_dimensions[chr(65 + i)].width = w
    for r in range(5, 105):
        for c in range(1, 13):
            cell = ws.cell(r, c)
            style_input(cell, {7: USD, 8: "yyyy-mm-dd", 9: "yyyy-mm-dd", 10: "yyyy-mm-dd"}.get(c))
    dv = DataValidation(type="list", formula1='"Lead,Quoted,Active,Paused,Lost"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("K5:K104")
    dv2 = DataValidation(type="list", formula1='"One-time,Weekly,Biweekly,Monthly,Nightly"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("F5:F104")
    if example:
        rows = [
            ("Sample Client A", "555-0101", "a@example.com", "Oak St", "Standard clean", "Biweekly", 130,
             date(2026, 8, 1), date(2026, 9, 12), date(2026, 9, 26), "Active", "Has two dogs"),
            ("Sample Client B", "555-0102", "b@example.com", "Pine Ave", "Move-out deep clean", "One-time", 410,
             date(2026, 9, 19), date(2026, 9, 19), date(2026, 10, 19), "Quoted", "Ask for referral"),
        ]
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                ws.cell(5 + i, 1 + j, v)
    ws.freeze_panes = "B5"


TEMPLATES = [
    ("Quote follow-up (2 days after sending a quote)",
     "Hi [Client Name], this is [Your Name] from [Business]. I sent your cleaning quote for [Service] on [Day]. "
     "Do you have any questions, or would you like me to hold a time for you on [Date]? Happy to adjust the scope if that helps."),
    ("Quote follow-up (7 days, last touch)",
     "Hi [Client Name], I'm closing out quotes from last week and didn't want to lose touch. If you're still "
     "interested, I can fit you in on [Date]. If now isn't the right time, no problem at all - just let me know."),
    ("After the first clean",
     "Hi [Client Name], thanks for having us today! Was everything how you wanted it? If anything was missed, tell me "
     "and we'll make it right. If you'd like us to come back on a regular schedule, I can set up [Frequency] visits at [Price]."),
    ("Recurring offer",
     "Hi [Client Name], since we've cleaned for you once, I can offer a recurring rate of [Price] per visit for "
     "[Frequency] service. Want me to put you on the calendar starting [Date]?"),
    ("Asking for a review",
     "Hi [Client Name], thank you for choosing [Business]. If you were happy with the clean, an honest review at "
     "[Review Link] helps a small business like ours. Whatever your experience was, we'd value your feedback."),
    ("Lapsed client check-in (60+ days)",
     "Hi [Client Name], it's been a while since your last clean. Would you like to schedule one for [Date]? "
     "I have openings on [Days]."),
    ("Price update notice (give 30 days' notice)",
     "Hi [Client Name], to keep our quality and our team paid fairly, our price for your [Service] will change from "
     "[Old Price] to [New Price] starting [Date]. Thank you for being a valued client; call me with any questions."),
    ("Reschedule / cancellation reply",
     "Hi [Client Name], no problem - I can move your clean to [Date/Time]. Per our policy, cancellations within "
     "[24] hours of the visit may carry a fee of [Amount]. Which time works best?"),
]


def build_templates(ws):
    title(ws, "Follow-Up Message Templates",
          "Copy, replace the [brackets], send. Get consent before texting or emailing marketing messages and honor opt-outs.")
    header_row(ws, 4, ["When to use", "Message"])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 110
    for i, (when, msg) in enumerate(TEMPLATES):
        r = 5 + i
        a = ws.cell(r, 1, when)
        b = ws.cell(r, 2, msg)
        for c in (a, b):
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOX
        ws.row_dimensions[r].height = 62


def build_start(ws, tier):
    title(ws, "Cleaning Business Job Profit & Quote Calculator", "Start here")
    lines = [
        ("What this does", "Tells you what to charge so a job actually makes the profit you want, then shows whether the job did."),
        ("Step 1", "Go to the Settings sheet. Enter your labor cost per hour, overhead %, target margin, mileage cost, recurring discount and minimum price."),
        ("Step 2", "Go to the Quote sheet. Describe the job in the yellow cells. Quote the price in 'QUOTE THIS PRICE'."),
        ("Step 3", "See the Examples sheet for three fictional jobs (a recurring home, a move-out deep clean, a small office)."),
    ]
    if tier >= 2:
        lines += [
            ("Step 4", "After each job, add a row on the Job Tracker with your ACTUAL hours, supplies and miles. Check the Dashboard to see which jobs lost money."),
        ]
    if tier >= 3:
        lines += [
            ("Step 5", "Use the Client Tracker for leads and recurring clients, the Follow-Up Templates for messages, and the 12-Month Tracker to compare profit with your goal."),
        ]
    lines += [
        ("Colors", "Yellow = you type here. Gray = calculated, don't type over it. Blue = key results."),
        ("Margin vs markup", "Margin = profit / price. Markup = profit / cost. Example: a job costs $100. A 40% margin means price = 100 / (1 - 0.40) = $166.67, "
                             "profit $66.67. That is a 66.7% markup but a 40% margin. This sheet uses margin so your profit % is what you really keep."),
        ("Google Sheets", "Upload this file to Google Drive and open with Google Sheets, or use File > Import. All formulas are standard and work in both."),
        ("Not included", "No macros. No tax, payroll or accounting features. " + DISCLAIMER),
        ("Support", "This is a template. It can't know your costs. If a result looks wrong, check your Settings first."),
    ]
    for i, (k, v) in enumerate(lines):
        r = 4 + i
        a = ws.cell(r, 1, k)
        a.font = Font(bold=True, color=NAVY)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(r, 2, v)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 48 if len(v) > 110 else 30
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110


def build(tier, example):
    wb = Workbook()
    start = wb.active
    start.title = "Start Here"
    settings = wb.create_sheet("Settings")
    quote = wb.create_sheet("Quote")
    examples = wb.create_sheet("Examples")
    build_start(start, tier)
    build_settings(settings, example)
    build_quote(quote, example)
    build_examples(examples)
    if tier >= 2:
        tracker = wb.create_sheet("Job Tracker")
        dash = wb.create_sheet("Dashboard")
        build_tracker(tracker, example)
        build_dashboard(dash)
    if tier >= 3:
        build_twelve_month(wb.create_sheet("12-Month Tracker"), example)
        build_clients(wb.create_sheet("Client Tracker"), example)
        build_templates(wb.create_sheet("Follow-Up Templates"))
    from openpyxl.worksheet.properties import PageSetupProperties
    for ws in wb:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return wb


def main():
    os.makedirs(OUT, exist_ok=True)
    names = {1: "Calculator", 2: "Calculator+Tracker", 3: "Bundle"}
    for tier in (1, 2, 3):
        for example in (False, True):
            wb = build(tier, example)
            fn = f"Cleaning-Quote-Profit-{names[tier]}-{'EXAMPLE' if example else 'BLANK'}.xlsx"
            wb.save(os.path.join(OUT, fn))
            print("wrote", fn)


if __name__ == "__main__":
    main()
