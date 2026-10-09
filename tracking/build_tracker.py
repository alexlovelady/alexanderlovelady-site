"""Builds revenue-tracker.xlsx: daily metrics per offer, fee-aware net, cash received, and a rent-goal summary."""
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

NAVY, INPUT, CALC = "1F3A4D", PatternFill("solid", fgColor="FFF6D6"), PatternFill("solid", fgColor="EEF2F4")
HEAD = PatternFill("solid", fgColor=NAVY)
thin = Side(style="thin", color="B8C2C8"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '"$"#,##0.00'
OFFERS = ["Fiverr: n8n/Make/Zapier fix", "Fiverr: Claude Code debugging", "Fiverr: Stripe fix", "Fiverr: Excel/CSV cleaning",
          "Cleaning Quote Calculator", "n8n pack: Agency Onboarding", "n8n pack: Cleaning Lead-to-Booking", "Other"]

def head(ws, row, labels):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row, i, l); c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 42

wb = Workbook()
# ---- Goal
g = wb.active; g.title = "Goal"
g["A1"] = "Rent goal"; g["A1"].font = Font(size=16, bold=True, color=NAVY)
rows = [("Amount needed NET (after fees), on top of other income", 2000, USD, True),
        ("Platform fee % (Fiverr reported 20%; check each platform)", 0.20, "0%", True),
        ("Gross sales needed", "=IF(B4>=1,0,B3/(1-B4))", USD, False),
        ("Net cash RECEIVED so far (payouts that have cleared)", "=SUMIFS(Orders!$K:$K,Orders!$L:$L,\"yes\")", USD, False),
        ("Net still pending clearance", "=SUMIFS(Orders!$K:$K,Orders!$L:$L,\"<>yes\")", USD, False),
        ("Remaining to reach goal (cleared only)", "=MAX(0,B3-B6)", USD, False),
        ("Goal date", date(2026, 11, 1), "yyyy-mm-dd", True),
        ("Days left", "=MAX(0,B9-TODAY())", "0", False)]
for i, (lab, v, fmt, inp) in enumerate(rows):
    r = 3 + i; g.cell(r, 1, lab).border = BOX; c = g.cell(r, 2, v); c.number_format = fmt; c.border = BOX
    c.fill = INPUT if inp else CALC
g.column_dimensions["A"].width = 62; g.column_dimensions["B"].width = 18
g["A12"] = "Cash only counts when the payout has cleared. Mark the 'Cleared?' column on the Orders sheet."
g["A12"].font = Font(italic=True, color="777777")

# ---- Orders
o = wb.create_sheet("Orders")
head(o, 1, ["Date ordered", "Offer", "Platform", "Customer / order #", "Gross price", "Fee %", "Fee $", "Hours spent", "Delivered on",
            "Cleared on", "Net received", "Cleared?", "Refunded $", "Net $/hour", "Notes"])
for r in range(2, 302):
    for c in (1, 2, 3, 4, 5, 6, 8, 9, 10, 12, 13, 15):
        cell = o.cell(r, c); cell.fill = INPUT; cell.border = BOX
    o.cell(r, 1).number_format = o.cell(r, 9).number_format = o.cell(r, 10).number_format = "yyyy-mm-dd"
    o.cell(r, 5).number_format = o.cell(r, 13).number_format = USD
    o.cell(r, 6).number_format = "0%"
    o.cell(r, 6, "=IF(E%d=\"\",\"\",Goal!$B$4)" % r)  # default fee, overwrite if platform differs
    o.cell(r, 7, f'=IF(E{r}="","",E{r}*F{r})'); o.cell(r, 11, f'=IF(E{r}="","",E{r}-G{r}-N(M{r}))')
    o.cell(r, 14, f'=IF(OR(E{r}="",H{r}="",H{r}=0),"",K{r}/H{r})')
    for c in (7, 11, 14): o.cell(r, c).fill = CALC; o.cell(r, c).border = BOX; o.cell(r, c).number_format = USD
dv = DataValidation(type="list", formula1="=Lists!$A$2:$A$20", allow_blank=True); o.add_data_validation(dv); dv.add("B2:B301")
dv2 = DataValidation(type="list", formula1='"Fiverr,Etsy,Gumroad,Direct,Other"', allow_blank=True); o.add_data_validation(dv2); dv2.add("C2:C301")
dv3 = DataValidation(type="list", formula1='"yes,no"', allow_blank=True); o.add_data_validation(dv3); dv3.add("L2:L301")
for col, w in zip("ABCDEFGHIJKLMNO", [12, 32, 11, 22, 12, 8, 10, 11, 12, 12, 13, 10, 11, 11, 30]): o.column_dimensions[col].width = w
o.freeze_panes = "C2"

# ---- Daily
d = wb.create_sheet("Daily Metrics")
head(d, 1, ["Date", "Offer", "Impressions", "Clicks", "Inquiries", "Orders", "Click rate", "Inquiry rate", "Order rate", "Notes / change made"])
start = date(2026, 10, 10)
for i in range(14 * len(OFFERS)):
    r = 2 + i; dt = start + timedelta(days=i // len(OFFERS))
    d.cell(r, 1, dt).number_format = "yyyy-mm-dd"; d.cell(r, 2, OFFERS[i % len(OFFERS)])
    for c in (3, 4, 5, 6, 10): d.cell(r, c).fill = INPUT
    d.cell(r, 7, f'=IF(N(C{r})=0,"",D{r}/C{r})'); d.cell(r, 8, f'=IF(N(D{r})=0,"",E{r}/D{r})'); d.cell(r, 9, f'=IF(N(E{r})=0,"",F{r}/E{r})')
    for c in (7, 8, 9): d.cell(r, c).number_format = "0.0%"; d.cell(r, c).fill = CALC
    for c in range(1, 11): d.cell(r, c).border = BOX
for col, w in zip("ABCDEFGHIJ", [12, 34, 12, 9, 10, 8, 11, 12, 11, 40]): d.column_dimensions[col].width = w
d.freeze_panes = "C2"

# ---- By Offer
b = wb.create_sheet("By Offer")
head(b, 1, ["Offer", "Impressions", "Clicks", "Inquiries", "Orders (logged)", "Gross sales", "Net received", "Hours", "Net $/hour", "Verdict"])
for i, name in enumerate(OFFERS):
    r = 2 + i; b.cell(r, 1, name)
    b.cell(r, 2, f"=SUMIFS('Daily Metrics'!$C:$C,'Daily Metrics'!$B:$B,$A{r})")
    b.cell(r, 3, f"=SUMIFS('Daily Metrics'!$D:$D,'Daily Metrics'!$B:$B,$A{r})")
    b.cell(r, 4, f"=SUMIFS('Daily Metrics'!$E:$E,'Daily Metrics'!$B:$B,$A{r})")
    b.cell(r, 5, f"=COUNTIFS(Orders!$B:$B,$A{r},Orders!$E:$E,\">0\")")
    b.cell(r, 6, f"=SUMIFS(Orders!$E:$E,Orders!$B:$B,$A{r})")
    b.cell(r, 7, f"=SUMIFS(Orders!$K:$K,Orders!$B:$B,$A{r})")
    b.cell(r, 8, f"=SUMIFS(Orders!$H:$H,Orders!$B:$B,$A{r})")
    b.cell(r, 9, f'=IF(H{r}=0,"",G{r}/H{r})')
    b.cell(r, 10, f'=IF(B{r}<100,"Too early (under 100 impressions)",IF(C{r}/B{r}<0.01,"Fix title/thumbnail (low clicks)",IF(D{r}=0,"Fix demo/price (clicks, no inquiries)",IF(E{r}=0,"Fix offer/response (inquiries, no orders)","Working - double down"))))')
    for c in (6, 7, 9): b.cell(r, c).number_format = USD
    for c in range(1, 11): b.cell(r, c).border = BOX
r = 2 + len(OFFERS); b.cell(r, 1, "TOTAL").font = Font(bold=True)
for c in (2, 3, 4, 5, 6, 7, 8):
    L = "ABCDEFGHIJ"[c - 1]; b.cell(r, c, f"=SUM({L}2:{L}{r-1})").font = Font(bold=True)
b.cell(r, 6).number_format = b.cell(r, 7).number_format = USD
for col, w in zip("ABCDEFGHIJ", [34, 12, 9, 10, 14, 13, 13, 8, 11, 40]): b.column_dimensions[col].width = w

l = wb.create_sheet("Lists"); l["A1"] = "Offers"; l["A1"].font = Font(bold=True)
for i, n in enumerate(OFFERS): l.cell(2 + i, 1, n)
l.column_dimensions["A"].width = 36
wb.save("revenue-tracker.xlsx"); print("saved")
